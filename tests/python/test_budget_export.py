"""Exporting the budget (ADR 0028 step 4): the lines in the panel's order,
CSV with its currency's decimal separator and the field separator the
spreadsheet expects, an Excel workbook with numbers in currency formats
(openpyxl), BUDGETEXPORT (EXPORTARPRESUPUESTO) and the panel's Export…."""

import csv
import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QFileDialog, QWidget
from pytestqt.qtbot import QtBot
from test_blocks import Harness
from test_budget_panel import EUR, USD, furnish
from ui_fakes import QuietUi

from arqstudiocad.budget import budget_of
from arqstudiocad.budget_export import (
    ExcelUnavailableError,
    budget_lines,
    export_budget,
    export_xlsx,
    number_format,
)
from arqstudiocad.core import CadApplication, CadCurrency, CadDocument
from arqstudiocad.events import CommandCancelled
from arqstudiocad.i18n import set_language
from arqstudiocad.ui.main_window import MainWindow
from arqstudiocad.ui.qt_services import QtUiServices


def furnished() -> CadDocument:
    document = CadApplication().create_document("Casa")
    furnish(document)
    return document


def read_csv(path: Path, delimiter: str) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.reader(file, delimiter=delimiter))


def test_the_lines_follow_the_panel() -> None:
    lines = budget_lines(budget_of(furnished()))
    assert [(line.kind, line.item, line.amount, line.currency) for line in lines] == [
        ("category", "Mobiliario", None, ""),
        ("item", "Asiento", "1500.00", "EUR"),
        ("item", "Mesa", "241.00", "USD"),
        ("item", "Silla", "450.00", "USD"),
        ("subtotal", "Subtotal Mobiliario", "691.00", "USD"),
        ("subtotal", "Subtotal Mobiliario", "1500.00", "EUR"),
        ("category", "No category", None, ""),
        ("item", "LAMPARA", "240.00", "EUR"),
        ("subtotal", "Subtotal No category", "240.00", "EUR"),
        ("without_cost", "Without cost (3)", None, ""),
        ("item", "Inodoro", None, ""),
        ("item", "JUEGO", None, ""),
        ("item", "PUERTA", None, ""),
        ("total", "Total USD", "691.00", "USD"),
        ("total", "Total EUR", "1740.00", "EUR"),
    ]
    chair = lines[3]
    assert (chair.material, chair.supplier, chair.unit, chair.quantity, chair.cost) == (
        "Madera",
        "Muebles SA",
        "each",
        9,
        "50.00",
    )


def test_csv_writes_each_currencys_decimals_without_thousands(tmp_path: Path) -> None:
    path = tmp_path / "presupuesto.csv"
    export_budget(budget_of(furnished()), path)
    assert path.read_bytes().startswith(b"\xef\xbb\xbf"), "a byte order mark, for accents"
    rows = read_csv(path, ",")
    assert rows[0] == [
        "Item",
        "Material",
        "Supplier",
        "Unit",
        "Quantity",
        "Unit cost",
        "Amount",
        "Currency",
    ]
    assert rows[1] == ["Mobiliario", "", "", "", "", "", "", ""]
    assert rows[2] == ["Asiento", "", "", "each", "1", "1500,00", "1500,00", "EUR"], (
        "the euro's comma; no thousands separator"
    )
    assert rows[4] == ["Silla", "Madera", "Muebles SA", "each", "9", "50.00", "450.00", "USD"]
    assert rows[11] == ["Inodoro", "", "", "each", "1", "", "", ""]
    assert rows[-1] == ["Total EUR", "", "", "", "", "", "1740,00", "EUR"]
    assert '"1500,00"' in path.read_text(encoding="utf-8-sig"), "quoted: it holds the separator"


def test_csv_separates_fields_with_semicolons_where_the_decimal_is_a_comma(
    tmp_path: Path,
) -> None:
    document = furnished()
    document.set_currencies([USD, EUR], "USD", ".", ",")
    path = tmp_path / "presupuesto.CSV"
    export_budget(budget_of(document), path)
    rows = read_csv(path, ";")
    assert rows[4] == ["Silla", "Madera", "Muebles SA", "each", "9", "50,00", "450,00", "USD"]
    assert ";" in path.read_text(encoding="utf-8-sig").splitlines()[0]


def test_excel_writes_numbers_in_currency_formats(tmp_path: Path) -> None:
    path = tmp_path / "presupuesto.xlsx"
    export_budget(budget_of(furnished()), path)
    sheet = load_workbook(path).active
    assert sheet is not None and sheet.title == "Budget"
    assert [c.value for c in sheet[1]][:2] == ["Item", "Material"]
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2"
    chair = [c.value for c in sheet[5]]
    assert chair == ["Silla", "Madera", "Muebles SA", "each", 9, 50, 450, "USD"]
    assert sheet["G5"].number_format == '"$" #,##0.00'
    assert sheet["F3"].number_format == '"€" #,##0.00'
    assert not sheet["A5"].font.bold
    assert sheet["A6"].value == "Subtotal Mobiliario" and sheet["A6"].font.bold
    assert Decimal(str(sheet["G6"].value)) == Decimal("691")
    assert [c.value for c in sheet[12]][4:7] == [1, None, None], "Inodoro: no cost"
    assert sheet["A16"].value == "Total EUR" and sheet["G16"].value == 1740
    assert sheet.max_row == 16


def test_number_formats_follow_the_currency() -> None:
    yen = CadCurrency("JPY", "Japanese yen", "¥", 0, True)
    odd = CadCurrency("XQT", "Quoted", 'Q"', 2, True)
    document = CadApplication().create_document("a")
    document.set_currencies([USD, yen, odd], "USD")
    budget_register = budget_of(document).register
    assert number_format(budget_register, "JPY", "3000") == '"¥" #,##0'
    assert number_format(budget_register, "XQT", "1.50") == '"Q""" #,##0.00'
    assert number_format(budget_register, "ZZZ", "1.5") == '"ZZZ" #,##0.0', "no symbol: its code"


def test_without_openpyxl_excel_export_says_so(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    path = tmp_path / "presupuesto.xlsx"
    with pytest.raises(ExcelUnavailableError, match="needs openpyxl"):
        export_xlsx(budget_of(furnished()), path)
    assert not path.exists()
    export_budget(budget_of(furnished()), tmp_path / "presupuesto.csv")


def test_another_type_is_refused(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match=r"not \.txt"):
        export_budget(budget_of(furnished()), tmp_path / "presupuesto.txt")


def harness(tmp_path: Path, name: str = "presupuesto.xlsx") -> tuple[Harness, QuietUi]:
    ui = QuietUi(export_path=tmp_path / name)
    h = Harness(ui)
    furnish(h.document)
    return h, ui


def cancelled(h: Harness) -> list[str]:
    return [e.reason for e in h.events if isinstance(e, CommandCancelled)]


def test_budgetexport_writes_the_whole_drawing_or_the_selection(tmp_path: Path) -> None:
    h, ui = harness(tmp_path)
    h.run("BUDGETEXPORT")
    path = tmp_path / "presupuesto.xlsx"
    assert ui.asked[-1] == f"export {h.document.name} - Budget.xlsx"
    assert h.messages()[-1] == f"Budget exported to {path}."
    sheet = load_workbook(path).active
    assert sheet is not None and sheet["A16"].value == "Total EUR"
    # With objects selected, their budget only, as BLOCKDATA takes them.
    ui.export_path = tmp_path / "piso.csv"
    entities = h.document.get_entities()
    lamp = next(e.id for e in entities if dict(e.details).get("Block name") == "LAMPARA")
    h.session.select(h.document, [lamp])
    h.run("BUDGETEXPORT")
    rows = read_csv(tmp_path / "piso.csv", ",")
    assert [row[0] for row in rows] == [
        "Item",
        "No category",
        "LAMPARA",
        "Subtotal No category",
        "Total EUR",
    ]
    assert rows[2][4] == "1"


def test_budgetexport_cancelled_or_failing_writes_nothing(tmp_path: Path) -> None:
    h, ui = harness(tmp_path)
    ui.export_path = None
    h.run("BUDGETEXPORT")
    assert list(tmp_path.iterdir()) == [] and len(cancelled(h)) == 1
    ui.export_path = tmp_path / "presupuesto.pdf"
    h.run("BUDGETEXPORT")
    assert cancelled(h)[-1] == "A budget is exported to .xlsx or .csv, not .pdf."
    ui.export_path = tmp_path / "missing" / "presupuesto.csv"
    h.run("BUDGETEXPORT")
    assert cancelled(h)[-1].startswith("presupuesto.csv could not be written: ")
    assert list(tmp_path.iterdir()) == []


def test_budgetexport_says_when_excel_is_unavailable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    h, _ = harness(tmp_path)
    h.run("BUDGETEXPORT")
    assert "needs openpyxl" in cancelled(h)[-1]
    assert list(tmp_path.iterdir()) == []


def test_an_empty_budget_is_not_exported(tmp_path: Path) -> None:
    ui = QuietUi(export_path=tmp_path / "x.csv")
    h = Harness(ui)
    h.run("BUDGETEXPORT")
    assert cancelled(h) == ["The budget is empty: there is nothing to export."]
    assert not any(a.startswith("export") for a in ui.asked)


def test_without_an_interface_the_file_is_typed(tmp_path: Path) -> None:
    h = Harness()
    furnish(h.document)
    h.run("BUDGETEXPORT")
    assert h.prompt().text == "File to export to (.xlsx or .csv):"
    h.run(str(tmp_path / "presupuesto.csv"))
    assert read_csv(tmp_path / "presupuesto.csv", ",")[1] == ["Mobiliario"] + [""] * 7


@pytest.fixture
def window(qtbot: QtBot, tmp_path: Path) -> MainWindow:
    ui = QuietUi(export_path=tmp_path / "presupuesto.csv")
    w = MainWindow(ui=ui)
    ui.panel_for = w.show_panel
    qtbot.addWidget(w)
    w.resize(1100, 700)
    w.show()
    qtbot.waitExposed(w)
    return w


def test_the_panels_export_writes_what_it_shows(window: MainWindow, tmp_path: Path) -> None:
    window.runner.execute("BUDGET")
    panel = window.budget_panel
    assert not panel.export_button.isEnabled(), "no drawing"
    window.runner.execute("NEW")
    assert not panel.export_button.isEnabled(), "nothing to export"
    document = window.session.active
    assert document is not None
    furnish(document)
    panel.refresh()
    assert panel.export_button.isEnabled()
    entities = document.get_entities()
    lamp = next(e.id for e in entities if dict(e.details).get("Block name") == "LAMPARA")
    window.session.select(document, [lamp])
    panel.export_button.click()
    path = tmp_path / "presupuesto.csv"
    assert read_csv(path, ",")[1][0] == "Mobiliario", "the whole drawing, as shown"
    panel.selection_only.setChecked(True)
    panel.export_button.click()
    assert read_csv(path, ",")[1][0] == "No category", "the selection, as shown"


def test_the_export_in_spanish(tmp_path: Path) -> None:
    try:
        set_language("es")
        h, ui = harness(tmp_path, "presupuesto.csv")
        h.run("EXPORTARPRESUPUESTO")
        rows = read_csv(tmp_path / "presupuesto.csv", ",")
        assert rows[0] == [
            "Partida",
            "Material",
            "Proveedor",
            "Unidad",
            "Cantidad",
            "P. unitario",
            "Parcial",
            "Moneda",
        ]
        assert rows[2][3] == "und" and rows[-1][0] == "Total EUR"
        assert ui.asked[-1] == f"export {h.document.name} - Presupuesto.xlsx"
    finally:
        set_language("en")


def test_the_save_dialog_gives_a_name_without_type_the_chosen_one(
    qtbot: QtBot, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    parent = QWidget()
    qtbot.addWidget(parent)
    services = QtUiServices(parent)
    answers: list[tuple[str, str]] = []
    asked: list[str] = []

    def dialog(_parent: QWidget, title: str, start: str, filters: str) -> tuple[str, str]:
        asked.append(f"{title} | {start} | {filters}")
        return answers.pop(0)

    monkeypatch.setattr(QFileDialog, "getSaveFileName", dialog)
    excel, comma = "Excel workbooks (*.xlsx)", "CSV files (*.csv)"
    base = str(tmp_path / "obra")
    answers += [(base, comma), (base, excel), (base + ".CSV", excel), ("", excel)]
    assert services.choose_budget_export_path("Casa - Budget.xlsx") == tmp_path / "obra.csv"
    assert asked[0] == f"Export budget | Casa - Budget.xlsx | {excel};;{comma}"
    assert services.choose_budget_export_path("x.xlsx") == tmp_path / "obra.xlsx"
    assert services.choose_budget_export_path("x.xlsx") == tmp_path / "obra.CSV", "typed: kept"
    assert asked[-1].split(" | ")[1] == str(tmp_path / "x.xlsx"), "starts in the last folder"
    assert services.choose_budget_export_path("x.xlsx") is None
