"""Exporting the budget (ADR 0028 D6): the budget as the Budget panel shows
it (groups, rows, subtotals, Without cost, totals) written to CSV or to an
Excel workbook, with a Currency column so every amount says its currency.

CSV is plain text, UTF-8 with a byte order mark so spreadsheets read
accents. Amounts are written with their currency's decimal separator and no
thousands separator; the field separator is a semicolon when the drawing's
decimal separator is a comma, as spreadsheets in those languages expect.

Excel needs openpyxl, imported only here and only when exporting: without
it, Excel export says so and CSV still works. Amounts are numbers with a
number format showing their currency's symbol and decimals; subtotals and
totals are values, not formulas.
"""

import csv
from collections.abc import Iterator
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Literal

from arqstudiocad.budget import Budget, BudgetRow
from arqstudiocad.currencies import CurrencyRegister
from arqstudiocad.i18n import tr

#: The file types the export writes, by suffix.
SUFFIXES = (".xlsx", ".csv")

Kind = Literal["category", "item", "subtotal", "without_cost", "total"]


class ExcelUnavailableError(Exception):
    """Excel export needs openpyxl, which is not installed."""


@dataclass(frozen=True)
class Line:
    """One line of the exported budget: a category heading, a block, a
    subtotal, the Without cost heading or a total. Amounts are decimal text
    in ``currency``."""

    kind: Kind
    item: str
    material: str = ""
    supplier: str = ""
    unit: str = ""
    quantity: int | None = None
    cost: str | None = None
    amount: str | None = None
    currency: str = ""


def headers() -> list[str]:
    """The columns, as the Budget panel names them, then Currency."""
    return [
        tr("Item"),
        tr("Material"),
        tr("Supplier"),
        tr("Unit"),
        tr("Quantity"),
        tr("Unit cost"),
        tr("Amount"),
        tr("Currency"),
    ]


def budget_lines(budget: Budget) -> list[Line]:
    """The budget's lines in the panel's order."""
    return list(_lines(budget))


def _lines(budget: Budget) -> Iterator[Line]:
    for group in budget.groups:
        title = group.category or tr("No category")
        yield Line("category", title)
        for row in group.rows:
            yield _item(row)
        for code, amount in group.subtotals.items():
            yield Line(
                "subtotal", tr("Subtotal {category}", category=title), amount=amount, currency=code
            )
    if budget.without_cost:
        yield Line("without_cost", tr("Without cost ({count})", count=len(budget.without_cost)))
        for row in budget.without_cost:
            yield _item(row)
    for code, amount in budget.totals.items():
        yield Line("total", tr("Total {currency}", currency=code), amount=amount, currency=code)


def _item(row: BudgetRow) -> Line:
    return Line(
        "item",
        row.item,
        row.material,
        row.supplier,
        tr("each"),
        row.quantity,
        row.cost,
        row.amount,
        row.currency or "",
    )


def export_budget(budget: Budget, path: Path) -> None:
    """Writes ``budget`` to ``path``, as CSV or Excel by its suffix. Raises
    ValueError for another suffix, ExcelUnavailableError without openpyxl,
    and OSError when the file cannot be written."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        export_csv(budget, path)
    elif suffix == ".xlsx":
        export_xlsx(budget, path)
    else:
        raise ValueError(tr("A budget is exported to .xlsx or .csv, not {suffix}.", suffix=suffix))


def export_csv(budget: Budget, path: Path) -> None:
    """Writes ``budget`` to ``path`` as CSV (see the module's docstring)."""
    register = budget.register
    delimiter = ";" if register is not None and register.decimal == "," else ","

    def number(amount: str | None, code: str) -> str:
        if amount is None or register is None:
            return ""
        return amount.replace(".", register.format_of(code)[1])

    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=delimiter)
        writer.writerow(headers())
        for line in budget_lines(budget):
            writer.writerow(
                [
                    line.item,
                    line.material,
                    line.supplier,
                    line.unit,
                    "" if line.quantity is None else line.quantity,
                    number(line.cost, line.currency),
                    number(line.amount, line.currency),
                    line.currency,
                ]
            )


def number_format(register: CurrencyRegister | None, code: str, amount: str) -> str:
    """An Excel number format showing currency ``code``'s symbol, or its
    code, and as many decimals as ``amount`` has. The spreadsheet shows
    the separators of its own language."""
    currency = register.get(code) if register is not None else None
    sign = currency.symbol if currency is not None and currency.symbol else code
    decimals = len(amount.partition(".")[2])
    digits = "#,##0." + "0" * decimals if decimals else "#,##0"
    quoted = sign.replace('"', '""')
    return f'"{quoted}" {digits}'


def export_xlsx(budget: Budget, path: Path) -> None:
    """Writes ``budget`` to ``path`` as an Excel workbook (see the
    module's docstring)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as error:
        raise ExcelUnavailableError(
            tr("Excel export needs openpyxl, which is not installed. Export to CSV instead.")
        ) from error
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    # Excel limits a sheet's name to 31 characters.
    sheet.title = tr("Budget")[:31]
    sheet.append(headers())
    bold = Font(bold=True)
    for cell in sheet[1]:
        cell.font = bold
    for line in budget_lines(budget):
        sheet.append(
            [
                line.item,
                line.material or None,
                line.supplier or None,
                line.unit or None,
                line.quantity,
                None if line.cost is None else Decimal(line.cost),
                None if line.amount is None else Decimal(line.amount),
                line.currency or None,
            ]
        )
        row = sheet.max_row
        for column, text in ((6, line.cost), (7, line.amount)):
            if text is not None:
                sheet.cell(row, column).number_format = number_format(
                    budget.register, line.currency, text
                )
        if line.kind != "item":
            for cell in sheet[row]:
                cell.font = bold
    for letter, width in zip("ABCDEFGH", (40, 18, 18, 8, 10, 16, 18, 10), strict=True):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)
